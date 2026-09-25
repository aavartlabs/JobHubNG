"""Phase 3: honest tailoring. The verifier must put back anything the model invents."""
from jobhub_poc import tailoring

RESUME = {
    "name": "Asha Rao", "headline": "Senior SRE", "location": "Bengaluru", "links": [],
    "summary": "SRE with 8 years running payment platforms.",
    "skills": ["Kubernetes", "Terraform", "AWS", "Python", "Prometheus"],
    "roles": [
        {"id": "r1", "title": "Senior SRE", "company": "Acme Pay", "location": "", "start": "Jan 2020", "end": "Present",
         "bullets": [
             {"id": "r1b1", "text": "Ran 12 Kubernetes clusters serving 3M daily payments."},
             {"id": "r1b2", "text": "Cut paging noise by 40% with Prometheus alert rules."},
             {"id": "r1b3", "text": "Organised the office football league."},
         ]},
        {"id": "r2", "title": "Systems Engineer", "company": "Beta Labs", "location": "", "start": "2016", "end": "Dec 2019",
         "bullets": [{"id": "r2b1", "text": "Automated server builds with Terraform on AWS."},
                     {"id": "r2b2", "text": "Wrote Python tooling for backups."}]},
    ],
    "education": ["B.E. Computer Science, 2015"],
}
JOB = {"title": "Platform Engineer", "company": "Zeta"}
REQS = {"required_skills": ["Kubernetes", "Go", "GCP"], "preferred_skills": ["Terraform"]}


def _raw(**over):
    raw = {
        "summary": "SRE with 8 years running payment platforms on Kubernetes and AWS.",
        "roles": [{"id": "r1", "bullets": [{"id": "r1b2", "text": "Cut paging noise by 40% using Prometheus alert rules."},
                                           {"id": "r1b1", "text": "Ran 12 Kubernetes clusters serving 3M daily payments."}]},
                  {"id": "r2", "bullets": [{"id": "r2b1", "text": "Automated server builds with Terraform on AWS."}]}],
        "skills": ["Kubernetes", "Terraform", "AWS"],
        "cover_note": "I run Kubernetes platforms at Acme Pay. I cut paging noise by 40%. I would enjoy the Platform Engineer role at Zeta.",
    }
    raw.update(over)
    return raw


def test_honest_rewording_and_selection_are_kept():
    out = tailoring.verify(_raw(), RESUME, JOB, REQS)
    r1 = out["roles"][0]
    assert [b["id"] for b in r1["bullets"]] == ["r1b2", "r1b1"]          # most relevant first
    assert r1["bullets"][0]["status"] == "reworded"
    assert r1["left_out"] == ["Organised the office football league."]
    assert out["summary_status"] == "tailored"
    assert out["skills"][:3] == ["Kubernetes", "Terraform", "AWS"] and set(out["skills"]) == set(RESUME["skills"])
    assert out["cover_note"].count(".") == 3 and out["log"] == []
    assert out["headline"] == "Senior SRE"                               # never retitled to the job


def test_a_new_number_a_new_name_or_a_job_skill_puts_the_original_back():
    raw = _raw(roles=[{"id": "r1", "bullets": [
        {"id": "r1b1", "text": "Ran 20 Kubernetes clusters serving 3M daily payments."},       # 12 -> 20
        {"id": "r1b2", "text": "Cut paging noise by 40% with Prometheus and Datadog."},        # new tool
        {"id": "r1b3", "text": "Organised the office football league and wrote Go services."},  # job skill
    ]}])
    out = tailoring.verify(raw, RESUME, JOB, REQS)
    b = {x["id"]: x for x in out["roles"][0]["bullets"]}
    assert all(x["status"] == "put_back" and x["text"] == x["original"] for x in b.values())
    assert "20" in b["r1b1"]["why"] and "Datadog" in b["r1b2"]["why"] and "Go" in b["r1b3"]["why"]
    assert len(out["log"]) == 3


def test_invented_ids_duplicates_and_roles_are_ignored_and_no_role_disappears():
    raw = _raw(roles=[{"id": "r9", "bullets": [{"id": "r9b1", "text": "Was CTO at Google."}]},
                      {"id": "r1", "bullets": [{"id": "r1b1", "text": "x"}, {"id": "r1b1", "text": "y"},
                                               {"id": "r2b1", "text": "stolen from another role"}]}])
    out = tailoring.verify(raw, RESUME, JOB, REQS)
    assert [r["id"] for r in out["roles"]] == ["r1", "r2"]
    assert [b["id"] for b in out["roles"][0]["bullets"]] == ["r1b1"]
    assert [b["id"] for b in out["roles"][1]["bullets"]] == ["r2b1", "r2b2"]   # skipped role: first bullets as written


def test_summary_skills_and_cover_note_cannot_claim_what_the_resume_doesnt():
    raw = _raw(summary="SRE with 10 years of GCP and Go experience.",
               skills=["Go", "GCP", "Python", "Kubernetes"],
               cover_note="I have run Kubernetes at Acme Pay. I have 5 years of GCP. I led a team at Google.")
    out = tailoring.verify(raw, RESUME, JOB, REQS)
    assert out["summary"] == RESUME["summary"] and out["summary_status"] == "original"
    assert out["skills"][:2] == ["Python", "Kubernetes"] and "Go" not in out["skills"]
    assert out["cover_note"] == "I have run Kubernetes at Acme Pay."
    assert any("Cover note" in line for line in out["log"])


def test_sentence_openers_are_grammar_not_names():
    assert tailoring.problems("Streamlined backups with Python tooling.", "Wrote Python tooling for backups.") == []
    assert tailoring.problems("Datadog dashboards for backups.", "Wrote Python tooling for backups.") == ["“Datadog”"]
    assert tailoring.problems("AWS builds.", "Automated builds.") == ["“AWS”"]


def test_the_review_form_is_the_users_own_words():
    out = tailoring.verify(_raw(), RESUME, JOB, REQS)
    form = {"summary": "My words.", "skills": "Go, Kubernetes, go", "cover_note": "Hi.", "contact": "a@b.c  +91 1",
            "role-0-bullets": "First\n\nSecond", "role-1-bullets": ""}
    edited = tailoring.from_form(form, out)
    assert edited["skills"] == ["Go", "Kubernetes"] and edited["contact"] == "a@b.c +91 1"
    assert [b["text"] for b in edited["roles"][0]["bullets"]] == ["First", "Second"]
    assert edited["roles"][1]["bullets"] == [] and edited["roles"][1]["title"] == "Systems Engineer"


# ---- the worker, the pages, downloads, the tracker ----

import io
import sqlite3
from datetime import datetime, timezone

import pytest
from cryptography.fernet import Fernet
from docx import Document

from jobhub_poc import config, crypto, db, job_requirements
from jobhub_poc.ai import llm, tasks, worker
from jobhub_poc.webapp.app import create_app

SESSION = "jobhub-auth.session_token"
GET_SESSION_URL = f"{config.AUTH_SERVICE_URL}/auth/get-session"


@pytest.fixture
def key(monkeypatch):
    monkeypatch.setattr(config, "RESUME_ENCRYPTION_KEY", Fernet.generate_key().decode())


def _seed_job(conn, jid=1):
    now = datetime.now(timezone.utc).isoformat()
    conn.execute("""INSERT INTO jobs (id, dedupe_key, source_site, title, company_name, location, apply_url,
                    description, is_remote, first_seen_at, last_seen_at, raw_json)
                    VALUES (?, ?, 's', 'Platform Engineer', 'Zeta', 'Bengaluru', 'https://e.example/a', 'x', 0, ?, ?, '{}')""",
                 (jid, f"k{jid}", now, now))
    job_requirements.store(conn, f"k{jid}", {**REQS, "min_years": None, "seniority": "unknown"}, "test")
    conn.commit()


def _store_resume(conn, user="u1"):
    now = datetime.now(timezone.utc).isoformat()
    conn.execute("""INSERT INTO resumes (owner_auth_user_id, filename, mime, size_bytes, file_enc, text_enc, structured_enc,
                    parse_status, consent_at, uploaded_at, updated_at) VALUES (?, 'r', 'm', 1, ?, ?, ?, 'done', ?, ?, ?)""",
                 (user, crypto.encrypt(b"f"), crypto.encrypt(b"t"), crypto.encrypt_json(RESUME), now, now, now))
    conn.commit()


def _client(conn, requests_mock, user="u1"):
    requests_mock.get(GET_SESSION_URL, json={"session": {}, "user": {"id": user, "email": f"{user}@x.com", "emailVerified": True}})
    app = create_app(test_conn=conn)
    app.config["TESTING"] = True
    c = app.test_client()
    c.set_cookie(SESSION, "x")
    return c


def _tailor(conn, client, monkeypatch, raw=None):
    client.post("/jobs/1/tailor")
    monkeypatch.setattr(llm, "generate", lambda *a, **k: raw or _raw())
    assert worker.run_once(conn)


def test_tailoring_is_queued_made_checked_and_stored_encrypted(conn, requests_mock, monkeypatch, key):
    _seed_job(conn); _store_resume(conn)
    c = _client(conn, requests_mock)
    assert "Tailor my resume for this job" in c.get("/jobs/1").get_data(as_text=True)
    c.post("/jobs/1/tailor")
    monkeypatch.setattr(llm, "available", lambda *a, **k: True)
    page = c.get("/jobs/1/tailor").get_data(as_text=True)
    assert "Tailoring your resume" in page and "data-match-pending" in page
    monkeypatch.setattr(llm, "available", lambda *a, **k: False)
    assert "back online" in c.get("/jobs/1/tailor").get_data(as_text=True)
    assert c.get("/jobs/1/tailor/status").status_code == 202

    bad = _raw(roles=[{"id": "r1", "bullets": [{"id": "r1b1", "text": "Ran 20 Kubernetes clusters on GCP."}]}])
    monkeypatch.setattr(llm, "generate", lambda *a, **k: bad)
    assert worker.run_once(conn)
    row = conn.execute("SELECT * FROM tailored_resumes").fetchone()
    assert row["status"] == "ready" and b"Kubernetes" not in row["data_enc"]
    assert c.get("/jobs/1/tailor/status").status_code == 200
    page = c.get("/jobs/1/tailor").get_data(as_text=True)
    assert "Built only from your checked resume" in page and "Your wording kept" in page
    assert "Ran 12 Kubernetes clusters serving 3M daily payments." in page and "20 Kubernetes" not in page
    assert "Review &amp; cover note" in c.get("/jobs/1").get_data(as_text=True)


def test_downloads_are_a_real_docx_and_a_print_page(conn, requests_mock, monkeypatch, key):
    _seed_job(conn); _store_resume(conn)
    c = _client(conn, requests_mock)
    _tailor(conn, c, monkeypatch)
    resp = c.get("/jobs/1/tailor/resume.docx")
    assert resp.status_code == 200 and 'filename="Asha-Rao-for-Zeta.docx"' in resp.headers["Content-Disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(resp.data)).paragraphs)
    assert "Asha Rao" in text and "Cut paging noise by 40%" in text and "football" not in text
    assert "u1@x.com" in text                                        # contact defaults to the account email
    printed = c.get("/jobs/1/tailor/print").get_data(as_text=True)
    assert "Save as PDF" in printed and "Acme Pay" in printed
    original = c.get("/profile/resume.docx")
    assert "football" in "\n".join(p.text for p in Document(io.BytesIO(original.data)).paragraphs)
    assert "Organised the office football league." in c.get("/profile/print").get_data(as_text=True)


def test_the_users_edits_are_saved_and_used(conn, requests_mock, monkeypatch, key):
    _seed_job(conn); _store_resume(conn)
    c = _client(conn, requests_mock)
    _tailor(conn, c, monkeypatch)
    c.post("/jobs/1/tailor/save", data={"summary": "Mine.", "skills": "AWS", "contact": "+91 90000 00000",
                                        "role-0-bullets": "My own bullet", "role-1-bullets": "", "cover_note": "Note."})
    assert conn.execute("SELECT status FROM tailored_resumes").fetchone()[0] == "edited"
    text = "\n".join(p.text for p in Document(io.BytesIO(c.get("/jobs/1/tailor/resume.docx").data)).paragraphs)
    assert "My own bullet" in text and "+91 90000 00000" in text and "Mine." in text


def test_tailored_resumes_are_the_owners_only_and_deleted_with_the_resume(conn, requests_mock, monkeypatch, key):
    _seed_job(conn); _store_resume(conn)
    _tailor(conn, _client(conn, requests_mock), monkeypatch)
    other = _client(conn, requests_mock, user="u2")
    assert "Upload it on your profile" in other.get("/jobs/1/tailor").get_data(as_text=True)
    assert other.get("/jobs/1/tailor/resume.docx").status_code == 302
    assert other.get("/profile/resume.docx").status_code == 302
    owner = _client(conn, requests_mock)
    owner.post("/profile/resume/delete")
    assert conn.execute("SELECT count(*) FROM tailored_resumes").fetchone()[0] == 0


def test_signed_out_visitors_are_sent_to_log_in(conn, requests_mock):
    _seed_job(conn)
    requests_mock.get(GET_SESSION_URL, json=None)
    client = create_app(test_conn=conn).test_client()
    for path in ("/jobs/1/tailor", "/jobs/1/tailor/resume.docx", "/profile/print"):
        assert client.get(path).status_code == 302
    assert client.post("/jobs/1/tailor").status_code == 302
    assert conn.execute("SELECT count(*) FROM ai_tasks").fetchone()[0] == 0


def test_the_tracker_asks_after_apply_and_my_jobs_filters(conn, requests_mock, key):
    _seed_job(conn); _seed_job(conn, 2)
    c = _client(conn, requests_mock)
    assert "Did you apply?" not in c.get("/jobs/1").get_data(as_text=True)
    c.get("/jobs/1/apply")
    assert "Did you apply?" in c.get("/jobs/1").get_data(as_text=True)
    c.post("/jobs/1/status", data={"status": "applied"})                 # saves it too
    c.post("/jobs/2/save")
    c.post("/jobs/1/status", data={"status": "made-up"})                 # ignored
    page = c.get("/jobs/1").get_data(as_text=True)
    assert "Did you apply?" not in page and 'value="applied" selected' in page
    mine = c.get("/saved").get_data(as_text=True)
    assert "All · 2" in mine and "Applied · 1" in mine and "Saved · 1" in mine
    only = c.get("/saved?status=applied").get_data(as_text=True)
    assert "/jobs/1\"" in only and "/jobs/2\"" not in only
    c.post("/saved/status", data={"key": "k1", "status": "interviewing"})
    assert conn.execute("SELECT status FROM saved_jobs WHERE job_dedupe_key = 'k1'").fetchone()[0] == "interviewing"


def test_old_databases_get_the_tracker_columns():
    old = sqlite3.connect(":memory:")
    old.row_factory = sqlite3.Row
    old.execute("""CREATE TABLE saved_jobs (owner_auth_user_id TEXT NOT NULL, job_dedupe_key TEXT NOT NULL, job_title TEXT,
                   job_company TEXT, job_location TEXT, job_apply_url TEXT, saved_at TEXT NOT NULL,
                   PRIMARY KEY (owner_auth_user_id, job_dedupe_key))""")
    old.execute("INSERT INTO saved_jobs VALUES ('u', 'k', 't', 'c', 'l', 'a', '2026-09-01')")
    db.init_db(old)
    db.init_db(old)
    assert dict(old.execute("SELECT status, status_at FROM saved_jobs").fetchone()) == {"status": "saved", "status_at": None}


def test_tailoring_waits_for_consent_to_changed_terms(conn, requests_mock, monkeypatch, key):
    _seed_job(conn); _store_resume(conn)
    conn.execute("UPDATE resumes SET consent_at = '2026-09-01T10:00:00+00:00'")
    conn.commit()
    monkeypatch.setattr(config, "AI_BACKEND", "openrouter")
    monkeypatch.setattr(config, "RESUME_CONSENT_SINCE", "2026-09-25T00:00:00+00:00")
    resp = _client(conn, requests_mock).post("/jobs/1/tailor")
    assert resp.status_code == 302 and "/profile" in resp.headers["Location"]
    assert conn.execute("SELECT COUNT(*) FROM ai_tasks").fetchone()[0] == 0
