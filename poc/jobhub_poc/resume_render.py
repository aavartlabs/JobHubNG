"""A resume (structured, or tailored by tailoring.py -- same shape) as a Word document. The
PDF route is the print page (templates/resume_print.html) and the browser's "Save as PDF":
no PDF engine on the Pi."""
import io
import re

from docx import Document
from docx.shared import Pt, RGBColor

_MUTED = RGBColor(0x55, 0x55, 0x55)


def dates(role):
    return " – ".join(x for x in (role.get("start"), role.get("end")) if x)


def filename(resume, suffix=""):
    """ASCII only: it goes in a Content-Disposition header."""
    def safe(text):
        return re.sub(r"[^A-Za-z0-9]+", "-", text or "").strip("-")[:60]
    base = safe(resume.get("name")) or "resume"
    return f"{base}-{safe(suffix)}.docx" if safe(suffix) else f"{base}.docx"


def to_docx(resume):
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(48)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)

    run = doc.add_paragraph().add_run(resume.get("name") or "")
    run.bold, run.font.size = True, Pt(18)
    line = " · ".join(x for x in (resume.get("headline"), resume.get("location"), resume.get("contact")) if x)
    if line:
        doc.add_paragraph().add_run(line).font.color.rgb = _MUTED
    for link in resume.get("links") or []:
        doc.add_paragraph().add_run(link).font.color.rgb = _MUTED

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run(text.upper())
        r.bold, r.font.size = True, Pt(10)

    if resume.get("summary"):
        heading("Summary")
        doc.add_paragraph(resume["summary"])
    if resume.get("skills"):
        heading("Skills")
        doc.add_paragraph(", ".join(resume["skills"]))
    if resume.get("roles"):
        heading("Experience")
        for role in resume["roles"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.add_run(role.get("title") or "").bold = True
            if role.get("company"):
                p.add_run(f" — {role['company']}")
            when = " · ".join(x for x in (dates(role), role.get("location")) if x)
            if when:
                r = p.add_run(f"\n{when}")
                r.font.color.rgb, r.font.size = _MUTED, Pt(9.5)
            for b in role.get("bullets") or []:
                doc.add_paragraph(b["text"], style="List Bullet")
    if resume.get("education"):
        heading("Education")
        for e in resume["education"]:
            doc.add_paragraph(e)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
