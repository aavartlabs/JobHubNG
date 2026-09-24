"""Plain text from an uploaded resume (PDF or DOCX), identified by its bytes, not its name."""
import io
import re
import zipfile


class ResumeUnreadable(ValueError):
    pass


def detect_kind(data: bytes):
    """"pdf", "docx" or None."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            return None
    return None


MIME = {"pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}


def extract_text(data: bytes, kind: str) -> str:
    try:
        if kind == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:10])
        elif kind == "docx":
            import docx
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            raise ResumeUnreadable("Upload a PDF or Word (.docx) file.")
    except ResumeUnreadable:
        raise
    except Exception as exc:  # noqa: BLE001 -- any parser failure is "we couldn't read it"
        raise ResumeUnreadable("We couldn't read that file. Try saving it again as PDF or DOCX.") from exc
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text).strip()
    if len(text) < 200:
        raise ResumeUnreadable("That file has almost no text we can read (a scanned image?). "
                               "Upload a PDF or DOCX with selectable text.")
    return text[:40000]
